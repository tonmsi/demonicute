"""Extract the received tribute workbook for source audit only; never save the XLSX."""
from __future__ import annotations

import hashlib
import json
from pathlib import Path
from datetime import date, datetime
from zipfile import ZipFile
from xml.etree import ElementTree as ET

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]
SOURCE = ROOT / "esempi di carte" / "DEMONICUTE 2026_SFIDE" / "TRIBUTI AGGIORNAMENTO SFIDE.xlsx"
OUT = Path(__file__).resolve().parent


def plain(value):
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    return value


def md(value):
    return str(value).replace("|", "\\|").replace("\r", "").replace("\n", "<br>")


source_hash = hashlib.sha256(SOURCE.read_bytes()).hexdigest()
book = load_workbook(SOURCE, data_only=False, read_only=False, keep_links=False)
cached = load_workbook(SOURCE, data_only=True, read_only=False, keep_links=False)
result = {
    "source": SOURCE.relative_to(ROOT).as_posix(),
    "sha256": source_hash,
    "sheet_names": book.sheetnames,
    "has_demoni_sheet": any(name.strip().casefold() == "demoni" for name in book.sheetnames),
    "authority": "New source under review; not a substitute for cutes 26-07-12.xlsx or author confirmation.",
    "active_game_content": False,
    "sheets": [],
}
with ZipFile(SOURCE) as archive:
    drawing_parts = [name for name in archive.namelist() if name.startswith("xl/drawings/") and name.endswith(".xml")]
    drawing_text = {name: [node.text for node in ET.fromstring(archive.read(name)).iter() if node.tag.endswith("}t") and node.text]
                    for name in drawing_parts}
    result["additional_objects"] = {
        "drawing_parts": drawing_parts,
        "drawing_text": drawing_text,
        "embedded_media": [name for name in archive.namelist() if name.startswith("xl/media/")],
    }
lines = ["# Tributi workbook extracted source", "", f"Source: `{result['source']}`", "", f"SHA-256: `{source_hash}`", "",
         "Read-only extraction. All nonempty cells, formulas and available cached values are retained with exact sheet/cell coordinates. "
         "The SFIDE folder name does not enable challenges or grant precedence over other sources. Original workbook unchanged.", ""]
for sheet in book:
    record = {
        "name": sheet.title, "state": sheet.sheet_state,
        "max_row": sheet.max_row, "max_column": sheet.max_column,
        "merged_ranges": [str(item) for item in sheet.merged_cells.ranges],
        "hidden_rows": [index for index, dimension in sheet.row_dimensions.items() if dimension.hidden],
        "hidden_columns": [key for key, dimension in sheet.column_dimensions.items() if dimension.hidden],
        "cells": [],
    }
    lines.extend([f"## {sheet.title}", "", f"State: `{sheet.sheet_state}`. Used bounds: {sheet.max_row} rows × {sheet.max_column} columns.", "",
                  "| Cell | Value or formula | Cached value | Comment |", "| --- | --- | --- | --- |"])
    for row in sheet.iter_rows():
        for cell in row:
            if cell.value is None and cell.comment is None:
                continue
            item = {"coordinate": cell.coordinate, "value": plain(cell.value), "data_type": cell.data_type}
            if cell.data_type == "f":
                item["cached_value"] = plain(cached[sheet.title][cell.coordinate].value)
            if cell.comment:
                item["comment"] = cell.comment.text
            if cell.hyperlink:
                item["hyperlink"] = cell.hyperlink.target or cell.hyperlink.location
            record["cells"].append(item)
            lines.append(f"| `{sheet.title}!{cell.coordinate}` | {md(cell.value) if cell.value is not None else ''} | {md(item.get('cached_value', ''))} | {md(item.get('comment', ''))} |")
    record["nonempty_cells"] = len(record["cells"])
    result["sheets"].append(record)
    lines.extend(["", f"Nonempty cells recorded: {record['nonempty_cells']}.", ""])
book.close()
cached.close()
result["formula_count"] = sum(cell["data_type"] == "f" for sheet in result["sheets"] for cell in sheet["cells"])
result["error_cells"] = [{"sheet": sheet["name"], "coordinate": cell["coordinate"], "value": cell["value"]}
                         for sheet in result["sheets"] for cell in sheet["cells"] if cell["data_type"] == "e"]
lines.extend(["## Extraction checks", "", f"Formula cells: {result['formula_count']}; stored error cells: {len(result['error_cells'])}.", "",
              "Additional drawing parts and text:", "", "```json", json.dumps(result["additional_objects"], ensure_ascii=False, indent=2), "```", ""])
if hashlib.sha256(SOURCE.read_bytes()).hexdigest() != source_hash:
    raise AssertionError("Source workbook changed during extraction")
(OUT / "tributes-extracted.json").write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
(OUT / "tributes-extracted.md").write_text("\n".join(lines), encoding="utf-8")
print(json.dumps({"sheet_names": result["sheet_names"], "has_demoni_sheet": result["has_demoni_sheet"],
                  "sheets": [{key: sheet[key] for key in ["name", "state", "max_row", "max_column", "nonempty_cells"]} for sheet in result["sheets"]]}, ensure_ascii=False, indent=2))
